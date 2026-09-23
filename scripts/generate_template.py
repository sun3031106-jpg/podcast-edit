"""
播客剪辑标注稿 docx 生成模板。
使用方法：在生成脚本中 import 这些辅助函数，传入分析结果生成标注稿。
"""

from docx import Document
from docx.shared import RGBColor, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


def create_doc(title_text):
    """创建标注稿文档，含标题和图例。"""
    doc = Document()

    style = doc.styles['Normal']
    font = style.font
    font.name = '微软雅黑'
    font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(title_text)
    run.bold = True
    run.font.size = Pt(16)

    legend = doc.add_paragraph()
    r = legend.add_run('标注说明：')
    r.bold = True
    r.font.size = Pt(12)

    p1 = doc.add_paragraph()
    r1 = p1.add_run('● 黑色文字 = 保留')
    r1.font.size = Pt(11)

    p2 = doc.add_paragraph()
    r2a = p2.add_run('● ')
    r2a.font.size = Pt(11)
    r2b = p2.add_run('【红色文字】')
    r2b.font.color.rgb = RGBColor(255, 0, 0)
    r2b.font.size = Pt(11)
    r2c = p2.add_run(' = 建议删除（啰嗦/重复/跑题/无效）')
    r2c.font.size = Pt(11)

    doc.add_paragraph()
    return doc


def add_header(doc, speaker, time):
    """添加发言人时间戳标头（蓝色加粗）。"""
    p = doc.add_paragraph()
    r = p.add_run(f'{speaker}   {time}')
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0, 0, 128)
    return p


def add_content(doc, segments):
    """
    添加正文段落。
    segments: list of (text, is_cut) — is_cut=True 的部分用红色【】标注。
    """
    p = doc.add_paragraph()
    for text, is_cut in segments:
        if is_cut:
            r = p.add_run(f'【{text}】')
            r.font.color.rgb = RGBColor(255, 0, 0)
        else:
            r = p.add_run(text)
        r.font.size = Pt(11)
    return p


def add_note(doc, text):
    """添加删除原因注释（灰色斜体）。"""
    p = doc.add_paragraph()
    r = p.add_run(f'    ▸ 删除原因：{text}')
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(150, 150, 150)
    r.italic = True
    return p
